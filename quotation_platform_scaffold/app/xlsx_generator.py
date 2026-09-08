"""
Generates a branded BHD quotation Excel workbook per account — the same
category-grouped structure as docx_generator.py (and the source ADC/Consilio
BOQ sheets), just as .xlsx instead of .docx. Same inputs, same totals, same
branding; pick whichever file format the client wants.
"""
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db_models import Account, DocTemplate, Project, Quote, TermsTemplate
from app.pricing_engine import calculate_quote_totals, group_line_items_by_category

BHD_FORMAT = '"BHD" #,##0.000'


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def generate_quote_spreadsheet(
    account: Account,
    doc_template: DocTemplate,
    project: Project,
    quote: Quote,
    terms_template: TermsTemplate | None,
    output_path: str,
) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"

    primary_hex = _hex(doc_template.primary_color)
    accent_hex = _hex(doc_template.accent_color)
    primary_fill = PatternFill(start_color=primary_hex, end_color=primary_hex, fill_type="solid")
    category_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
    white_bold = Font(color="FFFFFF", bold=True)

    show_tag = doc_template.show_tag_column
    columns = ["Sr.No."] + (["Tag"] if show_tag else []) + ["Items Description", "Unit", "QTY", "TOTAL PRICE BD"]
    n_cols = len(columns)

    row = 1
    if doc_template.logo_asset_path:
        try:
            img = XLImage(doc_template.logo_asset_path)
            img.width, img.height = 150, 60
            ws.add_image(img, "A1")
            row = 5  # leave room for the image before writing text
        except Exception:
            pass  # missing/corrupt asset shouldn't block generation

    ws.cell(row=row, column=1, value=account.firm_name).font = Font(size=16, bold=True, color=primary_hex)
    row += 1
    ws.cell(row=row, column=1, value="PROJECT QUOTATION").font = Font(size=11, bold=True, color=accent_hex)
    row += 2

    for label, value in [
        ("To:", quote.client_name),
        ("Project:", project.project_name),
        ("Project No.:", project.project_number or "-"),
        ("Quotation No.:", quote.quote_number or "-"),
    ]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    for i, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=i, value=col_name)
        cell.font = white_bold
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center")
    row += 1

    groups = group_line_items_by_category(quote.line_items)
    for group in groups:
        cat_cell = ws.cell(row=row, column=1, value=group["category"].name)
        cat_cell.font = Font(bold=True)
        cat_cell.fill = category_fill
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        row += 1

        for idx, item in enumerate(group["items"], start=1):
            col = 1
            ws.cell(row=row, column=col, value=idx)
            col += 1
            if show_tag:
                ws.cell(row=row, column=col, value=item.tag or "-")
                col += 1
            ws.cell(row=row, column=col, value=item.description)
            col += 1
            ws.cell(row=row, column=col, value=item.unit)
            col += 1
            ws.cell(row=row, column=col, value=float(item.quantity))
            col += 1
            price_cell = ws.cell(row=row, column=col, value=float(item.line_total))
            price_cell.number_format = BHD_FORMAT
            row += 1

    row += 1
    totals = calculate_quote_totals(quote.line_items, quote.vat_rate_snapshot, quote.payment_terms)
    label_col, value_col = n_cols - 1, n_cols

    def totals_row(label: str, value, bold: bool = False, color: str | None = None) -> None:
        nonlocal row
        label_cell = ws.cell(row=row, column=label_col, value=label)
        value_cell = ws.cell(row=row, column=value_col, value=float(value))
        value_cell.number_format = BHD_FORMAT
        if bold:
            label_cell.font = Font(bold=True)
            value_cell.font = Font(bold=True, color=color) if color else Font(bold=True)
        row += 1

    totals_row("Total", totals["subtotal"])
    totals_row(f"VAT ({float(quote.vat_rate_snapshot):g}%)", totals["vat_amount"])
    totals_row("Grand Total", totals["total"], bold=True, color=accent_hex)
    row += 1

    if totals["payment_schedule"]:
        ws.cell(row=row, column=1, value="Terms Of Payment:").font = Font(bold=True)
        row += 1
        for i, h in enumerate(["Percentage", "Description", "Amount"], start=1):
            ws.cell(row=row, column=i, value=h).font = Font(bold=True)
        row += 1
        for term in totals["payment_schedule"]:
            ws.cell(row=row, column=1, value=f"{term['pct']:g}%")
            ws.cell(row=row, column=2, value=term["label"])
            amount_cell = ws.cell(row=row, column=3, value=float(term["amount"]))
            amount_cell.number_format = BHD_FORMAT
            row += 1
        row += 1

    if terms_template and terms_template.sections:
        ws.cell(row=row, column=1, value="Terms & Conditions:").font = Font(bold=True)
        row += 1
        for section in terms_template.sections:
            if section.get("heading"):
                ws.cell(row=row, column=1, value=section["heading"]).font = Font(bold=True)
                row += 1
            ws.cell(row=row, column=1, value=section.get("body", ""))
            row += 1

    ws.column_dimensions["A"].width = 32
    for col_letter in (get_column_letter(i) for i in range(2, n_cols + 1)):
        ws.column_dimensions[col_letter].width = 18

    wb.save(output_path)
    return output_path
