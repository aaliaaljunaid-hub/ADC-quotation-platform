"""
Generates a branded BHD quotation Word document per account, mirroring the
ADC/Consilio sheets: grouped by category (categories.process_sequence order),
a bold category-header row, Sr.No reset per category, Unit/QTY/Total, then
Total / VAT / Grand Total, a payment schedule, terms & conditions, and a
signature block. Branding (logo, primary/accent colors, whether to show the
Tag column) comes from the account's DocTemplate.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.db_models import Account, DocTemplate, Project, Quote, TermsTemplate
from app.pricing_engine import calculate_quote_totals, group_line_items_by_category


def _hex_to_rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def _set_cell_shading(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shd)


def _bhd(amount) -> str:
    return f"BHD {amount:,.3f}"


def generate_quote_document(
    account: Account,
    doc_template: DocTemplate,
    project: Project,
    quote: Quote,
    terms_template: TermsTemplate | None,
    output_path: str,
) -> str:
    doc = Document()
    primary = doc_template.primary_color
    accent = doc_template.accent_color

    if doc_template.logo_asset_path:
        try:
            doc.add_picture(doc_template.logo_asset_path, width=Cm(4))
        except Exception:
            pass  # missing/corrupt asset shouldn't block document generation

    header = doc.add_paragraph()
    run = header.add_run(account.firm_name)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = _hex_to_rgb(primary)

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run("PROJECT QUOTATION")
    sub_run.font.size = Pt(11)
    sub_run.font.bold = True
    sub_run.font.color.rgb = _hex_to_rgb(accent)

    doc.add_paragraph()

    info = doc.add_paragraph()
    info.add_run("To: ").bold = True
    info.add_run(f"{quote.client_name}\n")
    info.add_run("Project: ").bold = True
    info.add_run(f"{project.project_name}\n")
    info.add_run("Project No.: ").bold = True
    info.add_run(f"{project.project_number or '-'}\n")
    info.add_run("Quotation No.: ").bold = True
    info.add_run(f"{quote.quote_number or '-'}")

    doc.add_paragraph()

    show_tag = doc_template.show_tag_column
    headers = ["Sr.No."] + (["Tag"] if show_tag else []) + ["Items Description", "Unit", "QTY", "TOTAL PRICE BD"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        _set_cell_shading(hdr_cells[i], primary)

    groups = group_line_items_by_category(quote.line_items)
    for group in groups:
        header_row_cells = table.add_row().cells
        merged = header_row_cells[0]
        for extra in header_row_cells[1:]:
            merged = merged.merge(extra)
        merged.text = group["category"].name
        merged.paragraphs[0].runs[0].font.bold = True
        _set_cell_shading(merged, "E8E8E8")

        for idx, item in enumerate(group["items"], start=1):
            row = table.add_row().cells
            col = 0
            row[col].text = str(idx)
            col += 1
            if show_tag:
                row[col].text = item.tag or "-"
                col += 1
            row[col].text = item.description
            col += 1
            row[col].text = item.unit
            col += 1
            row[col].text = f"{item.quantity:g}"
            col += 1
            row[col].text = _bhd(item.line_total)

    doc.add_paragraph()

    totals = calculate_quote_totals(quote.line_items, quote.vat_rate_snapshot, quote.payment_terms)

    totals_para = doc.add_paragraph()
    totals_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    totals_para.add_run(f"Total: {_bhd(totals['subtotal'])}\n")
    totals_para.add_run(f"VAT ({quote.vat_rate_snapshot:g}%): {_bhd(totals['vat_amount'])}\n")
    grand_run = totals_para.add_run(f"Grand Total: {_bhd(totals['total'])}")
    grand_run.font.bold = True
    grand_run.font.size = Pt(13)
    grand_run.font.color.rgb = _hex_to_rgb(accent)

    if totals["payment_schedule"]:
        doc.add_paragraph()
        doc.add_paragraph().add_run("Terms Of Payment:").bold = True
        pay_table = doc.add_table(rows=1, cols=3)
        hdr = pay_table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Percentage", "Description", "Amount"
        for cell in hdr:
            cell.paragraphs[0].runs[0].font.bold = True
        for term in totals["payment_schedule"]:
            row = pay_table.add_row().cells
            row[0].text = f"{term['pct']:g}%"
            row[1].text = term["label"]
            row[2].text = _bhd(term["amount"])

    if terms_template and terms_template.sections:
        doc.add_paragraph()
        doc.add_paragraph().add_run("Terms & Conditions:").bold = True
        for section in terms_template.sections:
            if section.get("heading"):
                doc.add_paragraph().add_run(section["heading"]).bold = True
            doc.add_paragraph(section.get("body", ""))

    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig.add_run("Name: _____________________________________     CPR: _______________________\n\n")
    sig.add_run("Signature: _____________________")

    doc.save(output_path)
    return output_path
